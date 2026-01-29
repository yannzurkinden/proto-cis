"""Report generation endpoints for PDF and Excel exports."""

import io
from datetime import date
from typing import Annotated, Optional

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import get_db
from app.dependencies import require_management, require_msp_or_above
from app.models.user import User
from app.repositories.beneficiary_repository import BeneficiaryRepository
from app.repositories.objective_repository import ObjectiveRepository
from app.repositories.skill_repository import SkillRepository
from app.repositories.time_tracking_repository import TimeTrackingRepository
from app.repositories.journal_repository import JournalRepository
from app.services.pdf import PDFService

router = APIRouter()


@router.get("/beneficiary/{beneficiary_id}/summary")
async def generate_beneficiary_summary(
    beneficiary_id: int,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_msp_or_above)],
):
    """Generate a PDF summary report for a beneficiary."""
    beneficiary_repo = BeneficiaryRepository(db)
    beneficiary = await beneficiary_repo.get_by_id_with_relations(beneficiary_id)

    if not beneficiary:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Beneficiary not found",
        )

    # Check access: MSP can only access their own beneficiaries
    if current_user.role == "MSP" and beneficiary.referent_id != current_user.id:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Access denied",
        )

    # Build data dict for PDF generation
    objective_repo = ObjectiveRepository(db)
    skill_repo = SkillRepository(db)
    time_repo = TimeTrackingRepository(db)
    journal_repo = JournalRepository(db)

    # Get objectives for this beneficiary
    objectives_list = await objective_repo.get_by_beneficiary(beneficiary_id)
    objectives_data = [
        {
            "titre": obj.title,
            "statut": obj.status,
            "progression": obj.progress or 0,
            "date_echeance": obj.due_date.isoformat() if obj.due_date else "-",
        }
        for obj in objectives_list
    ]

    # Get skill evaluations
    skills_list = await skill_repo.get_beneficiary_skills(beneficiary_id)
    skills_data = [
        {
            "competence": sk.skill.name if sk.skill else "-",
            "niveau": sk.level,
            "date_evaluation": sk.evaluation_date.isoformat() if sk.evaluation_date else "-",
        }
        for sk in skills_list
    ]

    # Get absences
    absences_list, _ = await time_repo.get_absences(
        beneficiary_id=beneficiary_id, skip=0, limit=100
    )
    absences_data = [
        {
            "date_debut": ab.start_date.isoformat() if ab.start_date else "-",
            "date_fin": ab.end_date.isoformat() if ab.end_date else "-",
            "motif": ab.absence_type,
            "justifie": ab.justification_document_id is not None,
        }
        for ab in absences_list
    ]

    # Get recent journal entries
    journal_entries = await journal_repo.get_recent_entries(
        beneficiary_id=beneficiary_id, limit=10
    )
    journal_data = [
        {
            "date": entry.entry_date.strftime("%d.%m.%Y") if entry.entry_date else "-",
            "auteur": entry.author.full_name if entry.author else "-",
            "contenu": entry.content or "",
        }
        for entry in journal_entries
    ]

    data = {
        "beneficiary": {
            "nom": beneficiary.last_name,
            "prenom": beneficiary.first_name,
            "date_naissance": beneficiary.date_of_birth.isoformat() if beneficiary.date_of_birth else "-",
            "numero_ai": beneficiary.ai_number or "-",
            "unite": beneficiary.unit.name if beneficiary.unit else "-",
            "msp_referent": beneficiary.referent.full_name if beneficiary.referent else "-",
        },
        "objectives": objectives_data,
        "skills": skills_data,
        "absences": absences_data,
        "journal_entries": journal_data,
    }

    pdf_service = PDFService()
    pdf_buffer = pdf_service.generate_beneficiary_summary(data)

    filename = f"summary_{beneficiary.last_name}_{beneficiary.first_name}.pdf"

    return StreamingResponse(
        io.BytesIO(pdf_buffer),
        media_type="application/pdf",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
        },
    )


@router.get("/activity")
async def generate_activity_report(
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_management)],
    unit_id: Optional[int] = None,
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
):
    """Generate an activity report (management only)."""
    import openpyxl

    # Use user's unit if not specified and user is RES
    if unit_id is None and current_user.role == "RES":
        unit_id = current_user.unit_id

    beneficiary_repo = BeneficiaryRepository(db)

    if unit_id:
        beneficiaries = await beneficiary_repo.get_by_unit(unit_id)
    else:
        beneficiaries, _ = await beneficiary_repo.get_all_filtered(
            skip=0, limit=10000
        )

    # Create Excel workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Activity Report"

    # Header row
    headers = [
        "ID",
        "Last Name",
        "First Name",
        "Status",
        "Unit",
        "Referent",
        "Entry Date",
        "Occupation Rate",
    ]
    ws.append(headers)

    # Style header
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = openpyxl.styles.Font(bold=True)

    # Data rows
    for b in beneficiaries:
        ws.append([
            b.id,
            b.last_name,
            b.first_name,
            b.status,
            b.unit.name if b.unit else "",
            b.referent.full_name if b.referent else "",
            b.entry_date.isoformat() if b.entry_date else "",
            float(b.occupation_rate) if b.occupation_rate else 0,
        ])

    # Auto-adjust column widths
    for column_cells in ws.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter
        for cell in column_cells:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except (TypeError, AttributeError):
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width

    # Save to buffer
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    today = date.today().isoformat()
    filename = f"activity_report_{today}.xlsx"

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
        },
    )


@router.get("/objectives")
async def export_objectives(
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_management)],
    unit_id: Optional[int] = None,
    status_filter: Optional[str] = Query(None, alias="status"),
):
    """Export objectives as an Excel file (management only)."""
    import openpyxl

    # Use user's unit if not specified and user is RES
    if unit_id is None and current_user.role == "RES":
        unit_id = current_user.unit_id

    objective_repo = ObjectiveRepository(db)

    objectives, total = await objective_repo.get_all_filtered(
        skip=0,
        limit=10000,
        unit_id=unit_id,
        status=status_filter,
    )

    # Create Excel workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Objectives"

    # Header row
    headers = [
        "ID",
        "Beneficiary",
        "Title",
        "Type",
        "Term",
        "Priority",
        "Status",
        "Progress (%)",
        "Due Date",
        "Created At",
    ]
    ws.append(headers)

    # Style header
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = openpyxl.styles.Font(bold=True)

    # Data rows
    for obj in objectives:
        ws.append([
            obj.id,
            obj.beneficiary.full_name if obj.beneficiary else "",
            obj.title,
            obj.objective_type or "",
            obj.term or "",
            obj.priority or "",
            obj.status,
            obj.progress or 0,
            obj.due_date.isoformat() if obj.due_date else "",
            obj.created_at.isoformat() if obj.created_at else "",
        ])

    # Auto-adjust column widths
    for column_cells in ws.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter
        for cell in column_cells:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except (TypeError, AttributeError):
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width

    # Save to buffer
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    today = date.today().isoformat()
    filename = f"objectives_export_{today}.xlsx"

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
        },
    )


@router.get("/absenteeism")
async def export_absenteeism(
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_management)],
    unit_id: Optional[int] = None,
    year: Optional[int] = None,
):
    """Export absence statistics as an Excel file (management only)."""
    import openpyxl

    # Use user's unit if not specified and user is RES
    if unit_id is None and current_user.role == "RES":
        unit_id = current_user.unit_id

    current_year = year or date.today().year

    beneficiary_repo = BeneficiaryRepository(db)
    time_repo = TimeTrackingRepository(db)

    if unit_id:
        beneficiaries = await beneficiary_repo.get_by_unit(unit_id)
    else:
        beneficiaries, _ = await beneficiary_repo.get_all_filtered(
            skip=0, limit=10000, status="active"
        )

    # Create Excel workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Absenteeism {current_year}"

    # Header row
    headers = [
        "Beneficiary",
        "Unit",
        "Total Days Absent",
        "Sick Days",
        "Vacation Days",
        "Accident Days",
        "Unauthorized Days",
        "Other Days",
        "Absence Rate (%)",
    ]
    ws.append(headers)

    # Style header
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = openpyxl.styles.Font(bold=True)

    # Data rows
    for b in beneficiaries:
        stats = await time_repo.get_absence_stats(
            beneficiary_id=b.id,
            year=current_year,
        )

        by_type = stats.by_type if hasattr(stats, "by_type") else {}

        ws.append([
            f"{b.last_name} {b.first_name}",
            b.unit.name if b.unit else "",
            stats.total_days if hasattr(stats, "total_days") else 0,
            by_type.get("sick", 0),
            by_type.get("vacation", 0),
            by_type.get("accident", 0),
            by_type.get("unauthorized", 0),
            by_type.get("other", 0),
            round(stats.absence_rate, 2) if hasattr(stats, "absence_rate") else 0,
        ])

    # Auto-adjust column widths
    for column_cells in ws.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter
        for cell in column_cells:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except (TypeError, AttributeError):
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width

    # Save to buffer
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    today = date.today().isoformat()
    filename = f"absenteeism_report_{current_year}_{today}.xlsx"

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
        },
    )
